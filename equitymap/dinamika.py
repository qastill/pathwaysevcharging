"""⚖️ Ekuitas vs Kesetaraan — lanjutan: struktur spasial, dinamika waktu, skenario kebijakan.

    python3 equitymap/prepare.py       # bila equity.js belum ada
    python3 equitymap/keadilan.py      # keadilan.json (kuintil provinsi dipakai di sini)
    python3 equitymap/dinamika.py      # -> equitymap/dinamika.json (digabung ke D.kd oleh keadilan_inject.py)
    python3 equitymap/keadilan_inject.py

Tiga pertanyaan lanjutan dari prinsip pertama:
  A. STRUKTUR SPASIAL — apakah ketimpangan mengelompok di ruang (Moran's I global, LISA per kabupaten:
     klaster tinggi-tinggi / rendah-rendah), atau tersebar acak? Bobot = ketetanggaan poligon geoBoundaries ADM2
     (queen, buffer ±0,5 km); pulau tanpa tetangga memakai 3 centroid terdekat. Uji permutasi 999×.
  B. DINAMIKA WAKTU — apakah ketimpangan membaik seiring jaringan tumbuh?
     (1) Jawa Barat: tanggal operasi nyata tiap unit charger (Master SPKLU Maret 2026, 636 unit).
     (2) Nasional PLN: master nasional tidak memuat tanggal; urutan ID SPKLU dipakai sebagai proksi urutan
         pembangunan, dikalibrasi ke tahun dengan 327 situs Jawa Barat yang tanggalnya diketahui (Spearman ≈0,90).
  C. SKENARIO KEBIJAKAN — berapa charger (dan rupiah) untuk menurunkan Gini ke 0,5/0,4/0,3/0,2 (water-filling:
     setiap kabupaten diangkat ke ambang per kapita yang sama, tanpa mengambil dari siapa pun), kabupaten mana
     yang menerima lebih dulu, dan batas trade-off cakupan vs pemerataan bila porsi anggaran f dialokasikan
     lewat aturan pemerataan (0 %, 25 %, 50 %, 75 %, 100 %) pada anggaran 250/500/1.000/2.000 situs baru.
"""
import csv, datetime, heapq, json, math, os, re

import h3
import numpy as np
import openpyxl
from scipy import sparse
from scipy.spatial import cKDTree
from scipy.stats import spearmanr
from shapely.geometry import shape
from shapely.strtree import STRtree

_trapz = getattr(np, "trapezoid", None) or getattr(np, "trapz")
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
os.chdir(ROOT)
rng = np.random.default_rng(42)

MASTER = "SPKLU_Indonesia_Lengkap_2026-06-08.xlsx"
MASTER_JB = "Master SPKLU Maret 2026.xlsx"
COVER_KM = 10.0
PERMS = 999
CH_SITE = 2            # charger per situs baru dalam simulasi
BUDGETS = [250, 500, 1000, 2000]
MIXES = [0.0, 0.25, 0.5, 0.75, 1.0]
GINI_TARGETS = [0.5, 0.4, 0.3, 0.2]
JAVA = {"Banten", "DKI Jakarta", "Jawa Barat", "Jawa Tengah", "DI Yogyakarta", "Jawa Timur"}

# ------------------------------------------------------------------ muat payload Peta Ekuitas
js = open(os.path.join(HERE, "equity.js"), encoding="utf-8").read()
E = json.loads(js[len("window.EQUITY="):].rstrip(";\n"))
ci = {c: i for i, c in enumerate(E["cols"])}
R6, KABS, PROVS = E["r6"], E["kabs"], E["provs"]
pi_name = [p["name"] for p in PROVS]
KS = [k for k in E["kab_stats"] if k["pop"] > 0]
n = len(KS)
pos = {k["idx"]: i for i, k in enumerate(KS)}            # idx KABS -> baris KS
pop_k = np.array([k["pop"] for k in KS], float)
ch_k = np.array([k["chargers"] for k in KS], float)
w10_k = np.array([k["within10"] for k in KS], float)
prov_k = np.array([k["prov"] for k in KS])
name_k = [k["name"] for k in KS]
is_kota = np.array([k["name"].startswith("Kota ") for k in KS])
is_java = np.array([pi_name[p] in JAVA for p in prov_k])
per100k = 1e5 * ch_k / pop_k
KD = json.load(open(os.path.join(HERE, "keadilan.json"), encoding="utf-8"))
Q1_PROVS = set(KD["quint"][0]["provs"])

H = [r[ci["h3"]] for r in R6]
hidx = {h: i for i, h in enumerate(H)}
pop_h = np.array([r[ci["pop"]] for r in R6], float)
kab_h = np.array([r[ci["kab"]] for r in R6])
kpos_h = np.array([pos.get(k, -1) for k in kab_h])
prov_h = np.array([KABS[k]["prov"] for k in kab_h])
d0 = np.array([r[ci["d_spklu"]] for r in R6], float)
ll = np.array([h3.cell_to_latlng(h) for h in H]); hlat, hlng = ll[:, 0], ll[:, 1]
P0 = pop_h.sum()
q1_h = np.array([pi_name[p] in Q1_PROVS for p in prov_h])


def xyz(lat, lng):
    """Koordinat 3-D di bola satuan × R — jarak tali busur ≈ haversine untuk radius ≤ 50 km."""
    la, lo = np.radians(np.asarray(lat, float)), np.radians(np.asarray(lng, float))
    R = 6371.0088
    return np.c_[R * np.cos(la) * np.cos(lo), R * np.cos(la) * np.sin(lo), R * np.sin(la)]


HX = xyz(hlat, hlng)


def gini(pop, val):
    pop, val = np.asarray(pop, float), np.asarray(val, float)
    o = np.argsort(np.divide(val, np.maximum(pop, 1e-9)))
    cp = np.concatenate([[0], np.cumsum(pop[o]) / pop.sum()])
    cv = np.concatenate([[0], np.cumsum(val[o]) / max(val.sum(), 1e-9)])
    return float(1 - 2 * _trapz(cv, cp))


def theil_between_pct(pop, val, group):
    pop, val, group = np.asarray(pop, float), np.asarray(val, float), np.asarray(group)
    P, Y = pop.sum(), val.sum()
    if Y <= 0:
        return 0.0, 0.0
    s = val > 0
    T = float(np.sum((val[s] / Y) * np.log((val[s] / Y) / (pop[s] / P))))
    b = 0.0
    for g in np.unique(group):
        m = group == g
        Pg, Yg = pop[m].sum(), val[m].sum()
        if Yg > 0:
            b += (Yg / Y) * math.log((Yg / Y) / (Pg / P))
    return round(T, 4), round(100 * b / T, 1) if T else 0.0


def within(sites_xyz, radius=COVER_KM, mask=None):
    """% penduduk heksagon (opsional: subset mask) yang berjarak ≤ radius dari situs mana pun."""
    if len(sites_xyz) == 0:
        return 0.0
    d, _ = cKDTree(sites_xyz).query(HX if mask is None else HX[mask])
    p = pop_h if mask is None else pop_h[mask]
    return round(100 * p[d <= radius].sum() / p.sum(), 2)


# ================================================================== A) STRUKTUR SPASIAL
print("A) struktur spasial …")
gj = json.load(open(os.path.join(HERE, "cache", "adm2.geojson"), encoding="utf-8"))
shp = {f["properties"]["shapeID"]: shape(f["geometry"]) for f in gj["features"]}
kabrows = list(csv.DictReader(open(os.path.join(HERE, "input", "kabupaten.csv"), encoding="utf-8")))
assert all(int(r["idx"]) == i for i, r in enumerate(kabrows))
geoms = [shp[kabrows[k["idx"]]["shapeID"]] for k in KS]
cent = [g.centroid for g in geoms]
clat = np.array([c.y for c in cent]); clng = np.array([c.x for c in cent])
gb = [g.buffer(0.005) for g in geoms]                      # ±0,5 km menutup celah poligon tersederhana
tree = STRtree(gb)
nbrs = []
for i, g in enumerate(gb):
    hits = [int(j) for j in tree.query(g, predicate="intersects") if j != i]
    nbrs.append(sorted(hits))
CX = xyz(clat, clng)
ctree = cKDTree(CX)
n_island = 0
for i in range(n):
    if not nbrs[i]:
        n_island += 1
        _, j = ctree.query(CX[i], k=4)
        nbrs[i] = [int(x) for x in j if x != i][:3]
deg = np.array([len(x) for x in nbrs])
rows_, cols_, vals_ = [], [], []
for i, nb in enumerate(nbrs):
    for j in nb:
        rows_.append(i); cols_.append(j); vals_.append(1.0 / len(nb))
W = sparse.csr_matrix((vals_, (rows_, cols_)), shape=(n, n))


def moran(x, Wm, perms=PERMS):
    x = np.asarray(x, float); z = x - x.mean(); m = len(z)
    denom = float(z @ z)
    lag = Wm @ z
    I = float(m / Wm.sum() * (z @ lag) / denom)
    sims = np.empty(perms)
    for p in range(perms):
        zp = rng.permutation(z)
        sims[p] = m / Wm.sum() * (zp @ (Wm @ zp)) / denom
    pval = (np.sum(sims >= I) + 1) / (perms + 1)
    return dict(I=round(I, 4), EI=round(-1 / (m - 1), 4), z=round(float((I - sims.mean()) / sims.std()), 2),
                p=round(float(pval), 4)), z, lag


def lisa(x, perms=PERMS):
    x = np.asarray(x, float); z = x - x.mean(); m2 = float(z @ z) / n
    Ii = np.zeros(n); lag = np.zeros(n); pv = np.ones(n)
    for i, nb in enumerate(nbrs):
        k = len(nb)
        lag[i] = z[nb].mean(); Ii[i] = z[i] * lag[i] / m2
        others = np.delete(z, i)
        samp = rng.choice(others, size=(perms, k))          # permutasi kondisional (dengan pengembalian)
        Is = z[i] * samp.mean(axis=1) / m2
        pv[i] = ((np.sum(Is >= Ii[i]) if Ii[i] >= 0 else np.sum(Is <= Ii[i])) + 1) / (perms + 1)
    cls = []
    for i in range(n):
        if pv[i] >= 0.05:
            cls.append("ns")
        elif z[i] >= 0 and lag[i] >= 0: cls.append("HH")
        elif z[i] < 0 and lag[i] < 0: cls.append("LL")
        elif z[i] >= 0: cls.append("HL")
        else: cls.append("LH")
    return Ii, lag, pv, cls, z


x_ch = np.log1p(per100k)                                    # log: skala per kapita sangat miring
mor_ch, z_ch, lag_ch = moran(x_ch, W)
mor_acc, z_acc, lag_acc = moran(w10_k, W)
Ii_ch, lagi_ch, p_ch, cls_ch, zz = lisa(x_ch)
Ii_acc, lagi_acc, p_acc, cls_acc, _ = lisa(w10_k)
CLS = ["HH", "LL", "HL", "LH", "ns"]


def cls_summary(cls):
    cls = np.array(cls)
    return [dict(cls=c, n=int((cls == c).sum()), pop=round(float(pop_k[cls == c].sum())),
                 pop_pct=round(100 * pop_k[cls == c].sum() / pop_k.sum(), 1),
                 chargers=int(ch_k[cls == c].sum()), per100k=round(1e5 * ch_k[cls == c].sum() / max(pop_k[cls == c].sum(), 1), 2),
                 within10=round(float(np.average(w10_k[cls == c], weights=pop_k[cls == c])), 1) if (cls == c).any() else 0)
            for c in CLS]


sum_ch, sum_acc = cls_summary(cls_ch), cls_summary(cls_acc)
cls_ch_a = np.array(cls_ch)
ll_prov = {}
for i in np.where(cls_ch_a == "LL")[0]:
    ll_prov[pi_name[prov_k[i]]] = ll_prov.get(pi_name[prov_k[i]], 0) + pop_k[i]
ll_prov = sorted(ll_prov.items(), key=lambda x: -x[1])[:10]
hh_list = sorted([i for i in range(n) if cls_ch[i] == "HH"], key=lambda i: -per100k[i])
ll_list = sorted([i for i in range(n) if cls_ch[i] == "LL"], key=lambda i: -pop_k[i])
zero_nb_zero = int(sum(1 for i in range(n) if ch_k[i] == 0 and all(ch_k[j] == 0 for j in nbrs[i])))
zero_total = int((ch_k == 0).sum())
# kabupaten tanpa charger yang seluruh tetangganya juga tanpa charger vs yang bertetangga dengan kabupaten ber-charger
# heksagon: Moran's I jarak ke SPKLU (cincin 1)
hn_rows, hn_cols = [], []
for i, h in enumerate(H):
    for nb in h3.grid_ring(h, 1):
        j = hidx.get(nb)
        if j is not None:
            hn_rows.append(i); hn_cols.append(j)
Wh = sparse.csr_matrix((np.ones(len(hn_rows)), (hn_rows, hn_cols)), shape=(len(H), len(H)))
rs = np.asarray(Wh.sum(axis=1)).ravel(); rs[rs == 0] = 1
Wh = sparse.diags(1 / rs) @ Wh
mor_hex, _, _ = moran(np.minimum(d0, 50), Wh, perms=199)
print(f"   Moran I charger/kapita (log) {mor_ch} · akses {mor_acc} · heksagon jarak {mor_hex} · pulau tanpa tetangga {n_island}")
print("   LISA charger:", [(s['cls'], s['n'], s['pop_pct']) for s in sum_ch])

# ================================================================== B) DINAMIKA WAKTU
print("B) dinamika waktu …")
# --- B1 Jawa Barat: tanggal operasi nyata
wb = openpyxl.load_workbook(MASTER_JB, read_only=True); ws = wb.worksheets[0]
it = ws.iter_rows(values_only=True); hdr = [str(c) for c in next(it)]; col = {c: i for i, c in enumerate(hdr)}
JB_PI = pi_name.index("Jawa Barat")
jb_names = {name_k[i]: i for i in range(n) if prov_k[i] == JB_PI}
cim_hex = h3.latlng_to_cell(-6.8722, 107.5425, 6)          # pusat Kota Cimahi (tidak punya heksagon sendiri)
cim_to = pos[kab_h[hidx[cim_hex]]] if cim_hex in hidx else jb_names["Bandung Barat"]


def jb_kab(s):
    s = re.sub(r"\s+", " ", str(s or "").strip().upper())
    if s.startswith("KAB"):
        nm = s.split(" ", 1)[1].strip(". ").title()
        return jb_names.get(nm)
    if s.startswith("KOTA"):
        nm = "Kota " + s.split(" ", 1)[1].strip(". ").title()
        if nm == "Kota Cimahi":
            return cim_to
        return jb_names.get(nm)
    return None


boxes = []
unmatched = set()
for r in it:
    try:
        lat, lng = float(r[col["LATITUDE"]]), float(r[col["LONGITUDE"]])
    except (TypeError, ValueError):
        continue
    t = r[col["TGL OPERASI"]]
    if not isinstance(t, datetime.datetime):
        continue
    k = jb_kab(r[col["KOTA/KAB"]])
    if k is None:
        unmatched.add(str(r[col["KOTA/KAB"]])); continue
    kw = float(re.sub(r"[^\d.]", "", str(r[col["KW"]] or "0").replace(",", ".")) or 0)
    boxes.append(dict(lat=lat, lng=lng, t=t, k=k, kw=kw, pln=str(r[col["MILIK"]] or "").upper() == "PLN"))
assert not unmatched, unmatched
jb_mask_k = prov_k == JB_PI
jb_mask_h = prov_h == JB_PI
jb_pop = pop_k[jb_mask_k]
cuts = [datetime.datetime(y, 12, 31) for y in range(2021, 2026)] + [max(b["t"] for b in boxes)]
jb_series = []
prev = 0
for c in cuts:
    sel = [b for b in boxes if b["t"] <= c]
    chk = np.zeros(n)
    for b in sel:
        chk[b["k"]] += 1
    sites = {(round(b["lat"], 4), round(b["lng"], 4)) for b in sel}
    sx = xyz([s[0] for s in sites], [s[1] for s in sites]) if sites else np.zeros((0, 3))
    cj = chk[jb_mask_k]
    kota_j, kab_j = is_kota[jb_mask_k], ~is_kota[jb_mask_k]
    T, tb = theil_between_pct(jb_pop, cj, kota_j.astype(int))
    new = [b for b in sel][prev:]
    jb_series.append(dict(
        label=c.strftime("%Y") if c.month == 12 else c.strftime("%b %Y"), units=len(sel), sites=len(sites), kw=round(sum(b["kw"] for b in sel)),
        gini=round(gini(jb_pop, cj), 3), zero=int((cj == 0).sum()),
        per100k=round(1e5 * cj.sum() / jb_pop.sum(), 2),
        kota_ratio=round((cj[kota_j].sum() / jb_pop[kota_j].sum()) / max(cj[kab_j].sum() / jb_pop[kab_j].sum(), 1e-9), 1),
        theil=T, theil_kota_pct=tb,
        within10=within(sx, 10, jb_mask_h), within5=within(sx, 5, jb_mask_h),
        new_units=len(sel) - prev, new_kota=int(sum(1 for b in sel[prev:] if is_kota[b["k"]])),
        new_pln=int(sum(1 for b in sel[prev:] if b["pln"])),
        top3_share=round(100 * np.sort(cj)[-3:].sum() / max(cj.sum(), 1), 1)))
    prev = len(sel)
print("   Jabar:", [(s["label"], s["units"], s["gini"], s["within10"]) for s in jb_series])

# --- B2 nasional PLN: urutan ID dikalibrasi ke tahun
wb = openpyxl.load_workbook(MASTER, read_only=True); ws = wb["Master SPKLU Indonesia"]
it = ws.iter_rows(values_only=True); hdr = [str(c) for c in next(it)]; col = {c: i for i, c in enumerate(hdr)}
spk = []
for r in it:
    if r[col["Latitude"]] in (None, "") or r[col["Longitude"]] in (None, ""):
        continue
    lat, lng = float(r[col["Latitude"]]), float(r[col["Longitude"]])
    if not (-11.5 < lat < 6.5 and 94 < lng < 142):
        continue
    st = str(r[col["Status"]] or "").strip().lower(); cat = str(r[col["Kategori"]] or "")
    active = (st in ("available", "inuse")) or (cat != "PLN" and st == "offline mode")
    sid = str(r[col["ID SPKLU"]])
    spk.append(dict(id=int(sid) if re.fullmatch(r"\d+", sid) else None, lat=lat, lng=lng,
                    ch=int(r[col["Jumlah Charger"]] or 0), active=active, pln=cat == "PLN"))
# kabupaten tiap situs = kabupaten heksagon res 6 tempat ia berada (fallback: heksagon berpenduduk terdekat)
htree = cKDTree(HX)
for s in spk:
    c6 = h3.latlng_to_cell(s["lat"], s["lng"], 6)
    j = hidx.get(c6)
    if j is None:
        _, j = htree.query(xyz([s["lat"]], [s["lng"]])[0]); j = int(j)
    s["kpos"] = kpos_h[j]
# pasangan kalibrasi: unit Jawa Barat ↔ situs PLN nasional dalam 300 m
stree = cKDTree(xyz([s["lat"] for s in spk], [s["lng"] for s in spk]))
pairs = {}
for b in boxes:
    d, j = stree.query(xyz([b["lat"]], [b["lng"]])[0])
    s = spk[j]
    if d < 0.3 and s["pln"] and s["id"] is not None and s["id"] < 10000:
        pairs[s["id"]] = min(pairs.get(s["id"], b["t"]), b["t"])
pid = np.array(sorted(pairs)); pyr = np.array([pairs[i].year for i in pid])
rho = spearmanr(pid, [pairs[i].toordinal() for i in pid]).statistic
YEARS = list(range(2021, 2027))
cuts_id = []
for y in YEARS[:-1]:
    m = (pyr == y) | (pyr == y + 1)
    ids, yrs = pid[m], pyr[m]
    best, bc = -1, None
    for c in np.unique(ids):
        acc = np.sum((ids < c) & (yrs == y)) + np.sum((ids >= c) & (yrs == y + 1))
        if acc > best:
            best, bc = acc, int(c)
    cuts_id.append(bc)


def year_of(i):
    for y, c in zip(YEARS, cuts_id):
        if i < c:
            return y
    return YEARS[-1]


pred = np.array([year_of(i) for i in pid])
calib = dict(n=len(pid), rho=round(float(rho), 3), acc=round(100 * np.mean(pred == pyr), 1),
             acc1=round(100 * np.mean(np.abs(pred - pyr) <= 1), 1), cuts=cuts_id,
             per_year=[dict(y=int(y), n=int((pyr == y).sum()), id_p10=int(np.percentile(pid[pyr == y], 10)),
                            id_med=int(np.median(pid[pyr == y])), id_p90=int(np.percentile(pid[pyr == y], 90))) for y in YEARS])
S_PLN = [s for s in spk if s["active"] and s["pln"]]
n_uncal = sum(1 for s in S_PLN if s["id"] is None or s["id"] >= 10000)
for s in S_PLN:
    s["year"] = YEARS[-1] if (s["id"] is None or s["id"] >= 10000) else year_of(s["id"])
S_MITRA = [s for s in spk if s["active"] and not s["pln"]]
nat_series = []
for y in YEARS:
    sel = [s for s in S_PLN if s["year"] <= y]
    chk = np.zeros(n)
    for s in sel:
        if s["kpos"] >= 0:
            chk[s["kpos"]] += s["ch"]
    sx = xyz([s["lat"] for s in sel], [s["lng"] for s in sel])
    T, tb = theil_between_pct(pop_k, chk, prov_k)
    kr = (chk[is_kota].sum() / pop_k[is_kota].sum()) / max(chk[~is_kota].sum() / pop_k[~is_kota].sum(), 1e-9)
    nat_series.append(dict(label=str(y) + (" (Jun)" if y == 2026 else ""), sites=len(sel), chargers=int(chk.sum()),
                           gini=round(gini(pop_k, chk), 3), gini_prov=round(gini(np.bincount(prov_k, pop_k, len(PROVS)), np.bincount(prov_k, chk, len(PROVS))), 3),
                           theil=T, theil_prov_pct=tb, zero=int((chk == 0).sum()), kota_ratio=round(kr, 1),
                           java_share=round(100 * chk[is_java].sum() / max(chk.sum(), 1), 1),
                           within10=within(sx), within25=within(sx, 25), q1_within10=within(sx, 10, q1_h),
                           provs=int(len(set(prov_k[chk > 0])))))
# titik akhir: + mitra non-PLN (tanpa tanggal)
sel = S_PLN + S_MITRA
chk = np.zeros(n)
for s in sel:
    if s["kpos"] >= 0:
        chk[s["kpos"]] += s["ch"]
sx = xyz([s["lat"] for s in sel], [s["lng"] for s in sel])
T, tb = theil_between_pct(pop_k, chk, prov_k)
kr = (chk[is_kota].sum() / pop_k[is_kota].sum()) / max(chk[~is_kota].sum() / pop_k[~is_kota].sum(), 1e-9)
nat_series.append(dict(label="2026 + mitra", sites=len(sel), chargers=int(chk.sum()), gini=round(gini(pop_k, chk), 3),
                       gini_prov=round(gini(np.bincount(prov_k, pop_k, len(PROVS)), np.bincount(prov_k, chk, len(PROVS))), 3),
                       theil=T, theil_prov_pct=tb, zero=int((chk == 0).sum()), kota_ratio=round(kr, 1),
                       java_share=round(100 * chk[is_java].sum() / chk.sum(), 1), within10=within(sx), within25=within(sx, 25),
                       q1_within10=within(sx, 10, q1_h), provs=int(len(set(prov_k[chk > 0])))))
print(f"   kalibrasi: n={calib['n']} rho={calib['rho']} akurasi tahun {calib['acc']}% (±1: {calib['acc1']}%) cuts={cuts_id} · PLN tanpa kalibrasi {n_uncal}")
print("   nasional:", [(s["label"], s["sites"], s["gini"], s["within10"]) for s in nat_series])

# ================================================================== C) SKENARIO KEBIJAKAN
print("C) skenario …")
# --- C1 water-filling: angkat semua kabupaten ke ambang t charger/100 rb
def wf(t):
    need = np.maximum(0, t * pop_k / 1e5 - ch_k)
    return float(need.sum()), gini(pop_k, ch_k + need), need


mean100k = 1e5 * ch_k.sum() / pop_k.sum()
grid = np.concatenate([[0], np.geomspace(0.05, 60, 80)])
wf_curve = [[round(float(t), 3), round(N), round(g, 4)] for t in grid for N, g, _ in [wf(t)]]
wf_targets = []
for tgt in GINI_TARGETS:
    lo, hi = 0.0, 500.0
    for _ in range(60):
        mid = (lo + hi) / 2
        if wf(mid)[1] > tgt: lo = mid
        else: hi = mid
    t = hi; N, g, need = wf(t)
    o = np.argsort(-need)
    provN = {}
    for i in range(n):
        if need[i] > 0:
            provN[pi_name[prov_k[i]]] = provN.get(pi_name[prov_k[i]], 0) + need[i]
    wf_targets.append(dict(
        target=tgt, t=round(t, 2), N=round(N), pct=round(100 * N / ch_k.sum(), 1), kab=int((need > 0).sum()),
        kota_share=round(100 * need[is_kota].sum() / N, 1), java_share=round(100 * need[is_java].sum() / N, 1),
        zero_share=round(100 * need[ch_k == 0].sum() / N, 1),
        prov=[dict(name=k, N=round(v), pct=round(100 * v / N, 1)) for k, v in sorted(provN.items(), key=lambda x: -x[1])[:8]],
        first=[dict(name=name_k[i], prov=pi_name[prov_k[i]], pop=round(float(pop_k[i])), chargers=int(ch_k[i]),
                    need=round(float(need[i]))) for i in o[:20]]))
print("   water-filling:", [(w["target"], w["t"], w["N"], w["kab"]) for w in wf_targets])

# --- C2 batas trade-off cakupan vs pemerataan (campuran dua aturan pada heksagon)
disk = [[hidx[m] for m in h3.grid_disk(h, 2) if m in hidx] for h in H]
hex_by_k = [[] for _ in range(n)]
for i, kp in enumerate(kpos_h):
    if kp >= 0:
        hex_by_k[kp].append(i)
cov0 = pop_h[d0 <= COVER_KM].sum()
q1_pop = pop_h[q1_h].sum()
EVERY = 25


def run_mix(f, B=max(BUDGETS)):
    covered = d0 <= COVER_KM
    chk = ch_k.copy()
    gain = np.array([sum(pop_h[j] for j in dk if not covered[j]) for dk in disk])
    heap = [(-g, i) for i, g in enumerate(gain) if g > 0]; heapq.heapify(heap)
    cov = cov0; n_eq = 0; picks = []; curve = []
    kab_sites = np.zeros(n)
    for s in range(1, B + 1):
        use_eq = n_eq < f * s - 1e-9
        if use_eq:
            n_eq += 1
            # kabupaten dengan charger per kapita terendah (seri: penduduk terbesar)
            k = int(np.lexsort((-pop_k, chk / pop_k))[0])
            cand = hex_by_k[k]
            best, bi = -1.0, cand[int(np.argmax(pop_h[cand]))]
            for i in cand:
                g = sum(pop_h[j] for j in disk[i] if not covered[j])
                if g > best:
                    best, bi = g, i
            i = bi
        else:
            i = None
            while heap:
                g, cand_i = heapq.heappop(heap)
                g2 = sum(pop_h[j] for j in disk[cand_i] if not covered[j])
                if g2 <= 0:
                    continue
                if heap and g2 < -heap[0][0] - 1e-9:
                    heapq.heappush(heap, (-g2, cand_i)); continue
                i = cand_i; break
            if i is None:                                   # semua sudah tercakup: sisa anggaran ikut aturan pemerataan
                k = int(np.lexsort((-pop_k, chk / pop_k))[0]); i = hex_by_k[k][int(np.argmax(pop_h[hex_by_k[k]]))]
        for j in disk[i]:
            if not covered[j]:
                covered[j] = True; cov += pop_h[j]
        kp = kpos_h[i]
        if kp >= 0:
            chk[kp] += CH_SITE; kab_sites[kp] += 1
        picks.append(i)
        if s % EVERY == 0 or s in BUDGETS:
            curve.append(dict(n=s, cov=round(100 * cov / P0, 2), gini=round(gini(pop_k, chk), 4),
                              q1=round(100 * pop_h[q1_h & covered].sum() / q1_pop, 2),
                              luar_java=round(100 * (1 - kab_sites[is_java].sum() / s), 1),
                              kota=round(100 * kab_sites[is_kota].sum() / s, 1)))
    return curve, kab_sites, picks


mix_runs = []
for f in MIXES:
    curve, kab_sites, picks = run_mix(f)
    at = {b: next(c for c in curve if c["n"] == b) for b in BUDGETS}
    top = np.argsort(-kab_sites)[:12]
    mix_runs.append(dict(f=f, curve=[c for c in curve if c["n"] % EVERY == 0 or c["n"] in BUDGETS],
                         at=[dict(n=b, **{k: v for k, v in at[b].items() if k != "n"}) for b in BUDGETS],
                         top_kab=[dict(name=name_k[i], prov=pi_name[prov_k[i]], sites=int(kab_sites[i])) for i in top if kab_sites[i] > 0]))
    print(f"   f={f}: ", [(a["n"], a["cov"], a["gini"], a["q1"]) for a in mix_runs[-1]["at"]])

# ================================================================== simpan
payload = dict(
    meta=dict(perms=PERMS, ch_site=CH_SITE, budgets=BUDGETS, mixes=MIXES, n_kab=n, n_island=n_island,
              deg_mean=round(float(deg.mean()), 1), cover_km=COVER_KM),
    spatial=dict(moran=dict(chargers=mor_ch, access=mor_acc, hex_dist=mor_hex),
                 summary=dict(chargers=sum_ch, access=sum_acc),
                 kab=[dict(name=name_k[i], prov=pi_name[prov_k[i]], lat=round(float(clat[i]), 3), lng=round(float(clng[i]), 3),
                           pop=round(float(pop_k[i])), per100k=round(float(per100k[i]), 2), within10=float(w10_k[i]),
                           z=round(float(zz[i]), 3), lag=round(float(lagi_ch[i]), 3), Ii=round(float(Ii_ch[i]), 3), p=round(float(p_ch[i]), 3),
                           cls=cls_ch[i], cls_acc=cls_acc[i], deg=int(deg[i])) for i in range(n)],
                 hh=[dict(name=name_k[i], prov=pi_name[prov_k[i]], per100k=round(float(per100k[i]), 2), lag=round(float(lagi_ch[i]), 2)) for i in hh_list[:15]],
                 ll=[dict(name=name_k[i], prov=pi_name[prov_k[i]], pop=round(float(pop_k[i])), chargers=int(ch_k[i]), within10=float(w10_k[i])) for i in ll_list[:15]],
                 ll_prov=[dict(name=k, pop=round(v)) for k, v in ll_prov],
                 zero_total=zero_total, zero_nb_zero=zero_nb_zero),
    time=dict(jabar=dict(series=jb_series, units=len(boxes), n_kab=int(jb_mask_k.sum()),
                         cimahi_to=name_k[cim_to], pln_units=int(sum(1 for b in boxes if b["pln"]))),
              nat=dict(series=nat_series, calib=calib, n_pln=len(S_PLN), n_uncal=n_uncal, n_mitra=len(S_MITRA))),
    scen=dict(mean100k=round(float(mean100k), 2), total=int(ch_k.sum()), wf_curve=wf_curve, wf_targets=wf_targets,
              mix=mix_runs, base=dict(cov=round(100 * cov0 / P0, 2), gini=round(gini(pop_k, ch_k), 4),
                                      q1=round(100 * pop_h[q1_h & (d0 <= COVER_KM)].sum() / q1_pop, 2))),
)
json.dump(payload, open(os.path.join(HERE, "dinamika.json"), "w", encoding="utf-8"), ensure_ascii=False, separators=(",", ":"))
print("→ dinamika.json %.0f KB" % (os.path.getsize(os.path.join(HERE, "dinamika.json")) / 1024))
